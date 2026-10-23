"""Builders for synthetic test documents (clearly fake content, TST- codes)."""
from __future__ import annotations

from pathlib import Path


def make_docx(path: Path, *, code="TST-FIN-POL-001", dept="Finance", version="1.1", status="Active",
              amount="₹ 1,00,000", extra_paragraphs: int = 0) -> Path:
    import docx

    d = docx.Document()
    d.add_paragraph("Fixture Expense Policy")
    header = d.add_table(rows=5, cols=2)
    header.style = "Table Grid"
    for i, (k, v) in enumerate([("Document ID", code), ("Department", dept), ("Version", version),
                                ("Status", status), ("Classification", "Internal")]):
        header.cell(i, 0).text, header.cell(i, 1).text = k, v
    d.add_paragraph("1. Purpose")
    d.add_paragraph(f"This fixture defines an approval limit of {amount} for testing purposes.")
    d.add_paragraph("5. Employee Expenses")
    d.add_paragraph("5.1 Business Travel")
    d.add_paragraph("Actual hotel cost is reimbursed up to the metro cap per night. Travel & lodging use code snake_case_code.")
    table = d.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    for r, (a, b) in enumerate([("Category", "Daily limit"), ("Meals", "₹1,200/day"), ("Local travel", "₹850/day")]):
        table.cell(r, 0).text, table.cell(r, 1).text = a, b
    for i in range(extra_paragraphs):
        d.add_paragraph(f"Additional paragraph {i} describing a reimbursement rule with limit ₹{(i + 1) * 100}.")
    d.save(path)
    return path


def make_xlsx(path: Path, *, code="TST-SAL-XLSX-002A", related="TST-SAL-POL-002") -> Path:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Metadata"
    ws.append(["Document ID", code])
    ws.append(["Department", "Sales"])
    ws.append(["Related Source Document", f"{related} — Fixture Returns Policy"])
    ws.append(["Status", "Active"])
    rules = wb.create_sheet("Approval Rules")
    rules.append(["Duration", "Approver", "SLA"])
    rules.append(["1-3 days", "Manager", "2 working days"])
    rules.append(["4-10 days", "Director", "5 working days"])
    wb.save(path)
    return path


def make_pdf(path: Path, pages: list[str]) -> Path:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=A4)
    for text in pages:
        y = 780
        for line in text.split("\n"):
            c.drawString(72, y, line)
            y -= 18
        c.showPage()
    c.save()
    return path


def build_corpus(raw: Path) -> Path:
    """Three departments with awkward names, mixed formats and three bad files."""
    (raw / "Finance").mkdir(parents=True)
    (raw / "Sales & Customer").mkdir()
    (raw / "Legal & Compliance" / "Contracts").mkdir(parents=True)
    make_docx(raw / "Finance" / "09_Fixture_Expense_Policy.docx", extra_paragraphs=3)
    (raw / "Finance" / "notes.txt").write_text("Plain note: reimbursement claims are filed within 30 days.", encoding="utf-8")
    make_xlsx(raw / "Sales & Customer" / "TST-SAL-XLSX-002A_Fixture_Matrix.xlsx")
    make_pdf(raw / "Sales & Customer" / "12_Fixture_Return_Policy.pdf",
             ["Fixture Return Policy\nReturns are accepted within 14 days of delivery.", "Refunds are processed within 7 working days."])
    (raw / "Legal & Compliance" / "Contracts" / "vendor_terms.md").write_text(
        "# Vendor Terms\n\n## 1. Payment\n\nVendors are paid within 45 days of invoice.\n", encoding="utf-8")
    # bad files
    (raw / "Finance" / "broken.docx").write_bytes(b"this is not a real docx")
    (raw / "Finance" / "empty.txt").write_bytes(b"")
    (raw / "Finance" / "drawing.dwg").write_bytes(b"CAD")
    return raw
