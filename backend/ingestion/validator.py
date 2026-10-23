"""Quality control for converted documents.

Docling output is *checked*, not trusted.  For every document we independently
extract the source text with a different library (python-docx, openpyxl,
pypdfium2), then measure how much of it survives in the Markdown:

* ``word_recall``   - share of distinct source words (>=3 letters) found in the Markdown;
* ``number_recall`` - same for numbers (amounts, dates, limits).  Dropped table cells
  show up here first, and numbers matter most in policy documents.

In addition: empty output, encoding damage, malformed tables, missing headings,
missing sheets/pages and table-count regressions are flagged.
"""
from __future__ import annotations

import logging
import re
from pathlib import Path

from backend.ingestion.markdown_utils import (
    PAGE_MARKER_RE,
    SHEET_HEADING_PREFIX,
    count_headings,
    find_encoding_problems,
)
from backend.models.documents import ParsedDocument, SourceDocument, ValidationResult

logger = logging.getLogger(__name__)

WORD_RECALL_WARN = 0.97
WORD_RECALL_ERROR = 0.85
NUMBER_RECALL_WARN = 0.98
NUMBER_RECALL_ERROR = 0.85

_WORD_RE = re.compile(r"[^\W\d_]{3,}", re.UNICODE)
_NUM_RE = re.compile(r"\d[\d,]*(?:\.\d+)?")


# ---------------------------------------------------------------------------- tokens
def word_set(text: str) -> set[str]:
    return {w.lower() for w in _WORD_RE.findall(text)}


def number_set(text: str) -> set[str]:
    out = set()
    for token in _NUM_RE.findall(text):
        token = token.replace(",", "")
        if "." in token:
            token = token.rstrip("0").rstrip(".")
        out.add(token)
    return out


# ---------------------------------------------------------------------------- reference extraction
class Reference:
    def __init__(self, text: str = "", tables: int | None = None, sheets: list[str] | None = None, pages: int | None = None):
        self.text = text
        self.tables = tables
        self.sheets = sheets or []
        self.pages = pages


def _ref_docx(path: Path) -> Reference:
    import docx  # python-docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            seen = set()
            for cell in row.cells:
                if id(cell._tc) in seen:  # merged cells repeat
                    continue
                seen.add(id(cell._tc))
                parts.append(cell.text)
    return Reference("\n".join(parts), tables=len(document.tables))


def _ref_xlsx(path: Path) -> Reference:
    import datetime as dt

    import openpyxl

    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    parts: list[str] = []
    sheets: list[str] = []
    for sheet in workbook.worksheets:
        has_values = False
        rows = []
        for row in sheet.iter_rows(values_only=True):
            for value in row:
                if value is None or isinstance(value, (dt.datetime, dt.date, dt.time)):
                    continue  # dates are rendered in locale-specific ways; not comparable
                rows.append(str(value))
                has_values = True
        if has_values:
            sheets.append(sheet.title)
            parts.append(sheet.title)
            parts.extend(rows)
    workbook.close()
    return Reference("\n".join(parts), sheets=sheets)


def _ref_pdf(path: Path) -> Reference:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(str(path))
    text = "\n".join(pdf[i].get_textpage().get_text_range() for i in range(len(pdf)))
    return Reference(text, pages=len(pdf))


def _ref_text(path: Path) -> Reference:
    return Reference(path.read_text(encoding="utf-8", errors="replace"))


def extract_reference(source: SourceDocument) -> Reference | None:
    """Independent extraction of the source text; None if the format has no reference extractor."""
    extractors = {"docx": _ref_docx, "xlsx": _ref_xlsx, "pdf": _ref_pdf, "txt": _ref_text, "md": _ref_text, "csv": _ref_text}
    extractor = extractors.get(source.file_type)
    if extractor is None:
        return None
    try:
        return extractor(source.path)
    except ImportError as exc:
        logger.warning("Reference extractor for %s unavailable: %s", source.file_type, exc)
    except Exception as exc:  # noqa: BLE001
        logger.warning("Could not build reference text for %s: %s", source.relative_path, exc)
    return None


# ---------------------------------------------------------------------------- markdown tables
def find_tables(markdown: str) -> list[list[str]]:
    """Return each Markdown table as a list of its lines."""
    tables, current = [], []
    for line in markdown.split("\n"):
        if line.lstrip().startswith("|"):
            current.append(line.strip())
        elif current:
            tables.append(current)
            current = []
    if current:
        tables.append(current)
    return tables


def _column_count(line: str) -> int:
    inner = line.strip()
    inner = inner[1:] if inner.startswith("|") else inner
    inner = inner[:-1] if inner.endswith("|") else inner
    return len(re.split(r"(?<!\\)\|", inner))


def table_problems(table: list[str]) -> list[str]:
    problems = []
    separator = re.compile(r"^\|?\s*:?-+:?\s*(\|\s*:?-+:?\s*)*\|?\s*$")
    if len(table) < 2 or not separator.match(table[1]):
        problems.append("missing header separator row")
    widths = {_column_count(line) for line in table}
    if len(widths) > 1:
        problems.append(f"inconsistent column counts {sorted(widths)}")
    header_cells = [c.strip() for c in table[0].strip("|").split("|")]
    if header_cells and all(not c for c in header_cells):
        problems.append("empty header row")
    return problems


# ---------------------------------------------------------------------------- main entry
def validate_markdown(
    markdown: str,
    source: SourceDocument,
    parsed: ParsedDocument | None = None,
    reference: Reference | None = None,
) -> ValidationResult:
    result = ValidationResult()
    body = markdown
    words_md = word_set(body)
    stats = {
        "chars": len(body),
        "words": len(re.findall(r"\S+", body)),
        "headings": count_headings(body),
    }
    tables = find_tables(body)
    stats["tables"] = len(tables)

    # 1. empty output ---------------------------------------------------------------
    if len(body.strip()) < 20:
        result.add("error", "EMPTY_OUTPUT", "Converted Markdown is empty or almost empty.")
        result.stats = stats
        return result

    # 2. encoding ---------------------------------------------------------------------
    problems = find_encoding_problems(body)
    if problems["replacement_chars"]:
        result.add("warning", "ENCODING_REPLACEMENT_CHARS", f"{problems['replacement_chars']} replacement character(s) (U+FFFD) found.")
    if problems["mojibake"]:
        result.add("warning", "ENCODING_MOJIBAKE", f"{problems['mojibake']} mojibake sequence(s) found (e.g. 'â‚¹' for '₹').")
    if problems["control_chars"]:
        result.add("warning", "ENCODING_CONTROL_CHARS", f"{problems['control_chars']} control character(s) found.")

    # 3. malformed tables ----------------------------------------------------------------
    for index, table in enumerate(tables, start=1):
        issues = table_problems(table)
        if issues:
            result.add("warning", "MALFORMED_TABLE", f"Table {index}: {', '.join(issues)}.")

    # 4. structure -------------------------------------------------------------------------
    if source.file_type in {"docx", "pdf"} and stats["chars"] > 1500 and stats["headings"] == 0:
        result.add("warning", "NO_HEADINGS", "No headings detected; chunking will fall back to paragraph boundaries.")

    if source.file_type == "pdf":
        page_markers = PAGE_MARKER_RE.findall(body)
        stats["pages_marked"] = len(page_markers)
        if parsed and parsed.page_count and len(page_markers) != parsed.page_count:
            result.add("warning", "PAGE_MARKER_MISMATCH", f"{len(page_markers)} page markers for {parsed.page_count} pages.")
        segments = PAGE_MARKER_RE.split(body)  # [pre, n1, text1, n2, text2, ...]
        for number, text in zip(segments[1::2], segments[2::2]):
            if len(text.strip()) < 20:
                result.add("warning", "EMPTY_PAGE", f"Page {number} has no extractable text (scanned page? enable DOCLING_DO_OCR).")

    # 5. comparison with the source ----------------------------------------------------------
    if reference is None:
        reference = extract_reference(source)
    if reference is None:
        result.add("info", "NO_REFERENCE", "No independent reference extractor for this file type; recall not measured.")
    else:
        src_words = word_set(reference.text)
        src_numbers = number_set(reference.text)
        if src_words:
            missing = sorted(src_words - words_md)
            recall = 1 - len(missing) / len(src_words)
            stats["word_recall"] = round(recall, 4)
            stats["source_words"] = len(src_words)
            if missing:
                stats["missing_words_sample"] = missing[:15]
            if recall < WORD_RECALL_ERROR:
                result.add("error", "LOW_TEXT_RECALL", f"Only {recall:.1%} of source words appear in the Markdown (possible missing content).")
            elif recall < WORD_RECALL_WARN:
                result.add("warning", "TEXT_RECALL", f"{recall:.1%} of source words appear in the Markdown; sample missing: {', '.join(missing[:8])}.")
        if src_numbers:
            md_numbers = number_set(body)
            missing_n = sorted(src_numbers - md_numbers)
            n_recall = 1 - len(missing_n) / len(src_numbers)
            stats["number_recall"] = round(n_recall, 4)
            if missing_n:
                stats["missing_numbers_sample"] = missing_n[:15]
            if n_recall < NUMBER_RECALL_ERROR:
                result.add("error", "LOW_NUMBER_RECALL", f"Only {n_recall:.1%} of source numbers appear in the Markdown.")
            elif n_recall < NUMBER_RECALL_WARN:
                result.add("warning", "NUMBER_RECALL", f"{n_recall:.1%} of source numbers appear in the Markdown; sample missing: {', '.join(missing_n[:8])}.")

        if reference.tables is not None:
            stats["source_tables"] = reference.tables
            if len(tables) < reference.tables:
                result.add("warning", "TABLE_COUNT", f"Source has {reference.tables} table(s) but Markdown has {len(tables)}.")
        if reference.sheets:
            found = re.findall(rf"(?m)^#{{1,6}}\s+{re.escape(SHEET_HEADING_PREFIX)}(.+)$", body)
            stats["sheets"] = len(found)
            absent = [s for s in reference.sheets if s not in found]
            if absent:
                result.add("warning", "MISSING_SHEETS", f"Sheet(s) not found in output: {', '.join(absent)}.")
        if reference.pages is not None:
            stats["source_pages"] = reference.pages

    result.stats = stats
    return result
